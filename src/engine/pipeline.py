"""Pipeline 调度器：按顺序执行排版规则"""

import copy
import os
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from src.engine.change_tracker import ChangeTracker, ChangeRecord
from src.engine.doc_tree import DocTree
from src.engine.rules.page_setup import PageSetupRule
from src.engine.rules.style_manager import StyleManagerRule
from src.engine.rules.heading_detect import HeadingDetectRule
from src.engine.rules.heading_numbering import HeadingNumberingRule
from src.engine.rules.caption_format import CaptionFormatRule
from src.engine.rules.section_format import SectionFormatRule
from src.engine.rules.table_format import TableFormatRule
from src.engine.rules.equation_table_format import EquationTableFormatRule
from src.engine.rules.whitespace_normalize import WhitespaceNormalizeRule
from src.engine.rules.citation_link import CitationLinkRule
from src.engine.rules.md_cleanup import MdCleanupRule
from src.engine.rules.toc_format import TocFormatRule
from src.engine.rules.header_footer import HeaderFooterRule
from src.engine.validation import ValidationRule
from src.scene.schema import SceneConfig, DEFAULT_PIPELINE_STEPS
from src.docx_io.sanitize import sanitize_docx
from src.docx_io.compare import generate_compare_doc
from src.docx_io.field_refresh import refresh_doc_fields_with_word


@dataclass
class PipelineResult:
    """Pipeline 执行结果"""
    success: bool
    status: str = "success"  # success | partial_success | failed | cancelled
    doc: object | None = None
    original_doc: object | None = None
    tracker: ChangeTracker = field(default_factory=ChangeTracker)
    output_paths: dict = field(default_factory=dict)
    failed_items: list[dict] = field(default_factory=list)
    error: str | None = None
    cancelled: bool = False

    def __post_init__(self) -> None:
        if self.cancelled:
            self.status = "cancelled"
            return
        if not self.success and self.status == "success":
            self.status = "failed"


# V0 可用规则注册表
RULE_REGISTRY = {
    "page_setup": PageSetupRule,
    "md_cleanup": MdCleanupRule,
    "style_manager": StyleManagerRule,
    "heading_detect": HeadingDetectRule,
    "heading_numbering": HeadingNumberingRule,
    "toc_format": TocFormatRule,
    "caption_format": CaptionFormatRule,
    "table_format": TableFormatRule,
    "equation_table_format": EquationTableFormatRule,
    "whitespace_normalize": WhitespaceNormalizeRule,
    "citation_link": CitationLinkRule,
    "section_format": SectionFormatRule,
    "header_footer": HeaderFooterRule,
    "validation": ValidationRule,
}

# V0 默认 pipeline 顺序（引用 schema.py 中的单一来源定义）
DEFAULT_PIPELINE = list(DEFAULT_PIPELINE_STEPS)


class Pipeline:
    """排版 Pipeline：加载文档 → 构建结构树 → 依次执行规则 → 输出结果"""

    def __init__(self, config: SceneConfig,
                 progress_callback=None,
                 cancel_requested=None):
        self.config = config
        self.tracker = ChangeTracker()
        self.progress_callback = progress_callback
        self.cancel_requested = cancel_requested

        # 按场景配置构建规则列表；仅在 None 时回退默认顺序。
        # 空列表表示显式禁用规则步骤，应被保留。
        step_names = DEFAULT_PIPELINE if config.pipeline is None else config.pipeline
        self.steps = []
        for name in step_names:
            rule_cls = RULE_REGISTRY.get(name)
            if rule_cls:
                self.steps.append(rule_cls())

    def _is_strict_mode_enabled(self) -> bool:
        """Strict mode can be configured by env var or scene config."""
        env_val = os.environ.get("DOCX_PIPELINE_STRICT_MODE")
        if env_val is not None and env_val.strip() != "":
            return env_val.strip().lower() in {"1", "true", "yes", "on"}
        return bool(getattr(self.config, "pipeline_strict_mode", True))

    def _critical_rules(self) -> set[str]:
        rules = getattr(self.config, "pipeline_critical_rules", None)
        if not isinstance(rules, list) or not rules:
            return set(DEFAULT_PIPELINE)
        return {str(x).strip() for x in rules if str(x).strip()}

    def _collect_failed_items(self) -> list[dict]:
        items: list[dict] = []
        for rec in self.tracker.get_failures():
            items.append({
                "rule_name": rec.rule_name,
                "target": rec.target,
                "section": rec.section,
                "change_type": rec.change_type,
                "paragraph_index": rec.paragraph_index,
                "reason": rec.failure_reason or "",
            })
        return items

    def _finalize_result(self, *, doc, original_doc, output_paths: dict) -> PipelineResult:
        """Build structured status for success/partial/failure states."""
        failed_items = self._collect_failed_items()
        if not failed_items:
            return PipelineResult(
                success=True,
                status="success",
                doc=doc,
                original_doc=original_doc,
                tracker=self.tracker,
                output_paths=output_paths,
                failed_items=[],
            )

        strict_mode = self._is_strict_mode_enabled()
        critical_rules = self._critical_rules()
        has_critical_failure = any(
            item.get("rule_name") in critical_rules for item in failed_items
        )

        if strict_mode and has_critical_failure:
            return PipelineResult(
                success=False,
                status="failed",
                doc=doc,
                original_doc=original_doc,
                tracker=self.tracker,
                output_paths=output_paths,
                failed_items=failed_items,
                error=f"关键步骤失败 {len(failed_items)} 项（严格模式）",
            )

        return PipelineResult(
            success=True,
            status="partial_success",
            doc=doc,
            original_doc=original_doc,
            tracker=self.tracker,
            output_paths=output_paths,
            failed_items=failed_items,
            error=f"存在未成功项 {len(failed_items)} 项",
        )

    def _emit_progress(self, current: int, total: int, message: str):
        if self.progress_callback:
            self.progress_callback(current, total, message)

    def _is_cancel_requested(self) -> bool:
        if self.cancel_requested is None:
            return False
        try:
            return bool(self.cancel_requested())
        except Exception:
            return False

    def _cancel_result(self, error: str = "用户已取消") -> PipelineResult:
        return PipelineResult(
            success=False,
            status="cancelled",
            tracker=self.tracker,
            failed_items=self._collect_failed_items(),
            error=error,
            cancelled=True,
        )

    def run(self, doc_path: str) -> PipelineResult:
        """执行完整排版流程"""
        if self._is_cancel_requested():
            return self._cancel_result()
        doc_path = Path(doc_path)
        if not doc_path.exists():
            return PipelineResult(success=False,
                                  status="failed",
                                  error=f"文件不存在: {doc_path}")

        disable_refresh = os.environ.get("DOCX_DISABLE_FIELD_REFRESH", "").strip() == "1"
        planned_refresh_step = bool(self.config.output.final_docx and not disable_refresh)
        # 阶段：加载、结构分析、规则执行、保存、(可选)域刷新、报告
        total_steps = len(self.steps) + 4 + (1 if planned_refresh_step else 0)
        step_idx = 0

        # 1. 预处理 & 加载文档
        self._emit_progress(step_idx, total_steps, "正在加载文档…")
        if self._is_cancel_requested():
            return self._cancel_result()
        try:
            sanitize_docx(str(doc_path))
        except Exception:
            pass  # 预处理失败不阻断，继续尝试加载
        if self._is_cancel_requested():
            return self._cancel_result()
        try:
            doc = Document(str(doc_path))
        except Exception as e:
            return PipelineResult(success=False,
                                  status="failed",
                                  error=f"无法打开文档: {e}")

        # 深拷贝原始文档用于对比稿
        original_doc = copy.deepcopy(doc)
        step_idx += 1
        if self._is_cancel_requested():
            return self._cancel_result()

        # 2. 构建文档结构树（根据 format_scope 决定自动/手动模式）
        self._emit_progress(step_idx, total_steps, "正在分析文档结构…")
        if self._is_cancel_requested():
            return self._cancel_result()
        doc_tree = DocTree()
        scope = self.config.format_scope
        body_start = None
        if scope.mode == "manual":
            if scope.body_start_index is not None:
                body_start = scope.body_start_index
            elif scope.body_start_page is not None:
                body_start = self._find_page_start_index(
                    doc, scope.body_start_page)
            elif scope.body_start_keyword:
                body_start = self._find_keyword_index(
                    doc, scope.body_start_keyword)
        doc_tree.build(doc, body_start_index=body_start)
        context = {
            "doc_tree": doc_tree,
            "format_scope": scope,
        }
        step_idx += 1
        if self._is_cancel_requested():
            return self._cancel_result()

        # 3. 依次执行规则
        for rule in self.steps:
            if self._is_cancel_requested():
                return self._cancel_result()
            self._emit_progress(step_idx, total_steps,
                                f"正在执行规则：{rule.description}")
            try:
                rule.apply(doc, self.config, self.tracker, context)
            except Exception as e:
                self.tracker.record(
                    rule_name=rule.name,
                    target="pipeline",
                    section="global",
                    change_type="error",
                    before="",
                    after="",
                    paragraph_index=-1,
                    success=False,
                    failure_reason=str(e),
                )
            step_idx += 1
            if self._is_cancel_requested():
                return self._cancel_result()

        # 4. 保存输出文件
        if self._is_cancel_requested():
            return self._cancel_result()
        self._emit_progress(step_idx, total_steps, "正在保存输出文件…")
        output_paths = self._save_outputs(doc, doc_path)
        step_idx += 1
        compare_path = output_paths.get("compare")
        if compare_path and original_doc is not None:
            try:
                # 序列化→反序列化终稿，避免 deepcopy 对 pipeline 修改后的
                # lxml 树产生损坏的克隆体（导致修订标记为空）。
                from io import BytesIO
                buf = BytesIO()
                doc.save(buf)
                buf.seek(0)
                final_doc_clean = Document(buf)

                generate_compare_doc(
                    original_doc=original_doc,
                    tracker=self.tracker,
                    output_path=compare_path,
                    final_doc=final_doc_clean,
                    include_text=bool(self.config.output.compare_text),
                    include_formatting=bool(self.config.output.compare_formatting),
                )
            except Exception as e:
                self.tracker.record(
                    rule_name="pipeline",
                    target="compare_doc",
                    section="global",
                    change_type="error",
                    before="",
                    after="",
                    paragraph_index=-1,
                    success=False,
                    failure_reason=f"对比稿生成失败: {e}",
                )

        if self._is_cancel_requested():
            return PipelineResult(
                success=False,
                status="cancelled",
                doc=doc,
                original_doc=original_doc,
                tracker=self.tracker,
                output_paths=output_paths,
                failed_items=self._collect_failed_items(),
                error="用户已取消",
                cancelled=True,
            )

        final_path = output_paths.get("final")
        if final_path and not disable_refresh:
            self._emit_progress(step_idx, total_steps, "正在更新目录与页码…")
            try:
                refresh_timeout = int(
                    os.environ.get("DOCX_FIELD_REFRESH_TIMEOUT_SEC", "10")
                )
            except ValueError:
                refresh_timeout = 10
            refresh_timeout = max(10, min(refresh_timeout, 600))
            ok, detail = refresh_doc_fields_with_word(
                final_path, timeout_sec=refresh_timeout
            )
            self.tracker.record(
                rule_name="pipeline",
                target="final_doc_fields",
                section="global",
                change_type="refresh",
                before="TOC/fields pending",
                after="TOC/fields refreshed" if ok else "TOC/fields refresh skipped",
                paragraph_index=-1,
                success=ok,
                failure_reason=None if ok else detail,
            )
            step_idx += 1
        elif final_path and disable_refresh:
            self.tracker.record(
                rule_name="pipeline",
                target="final_doc_fields",
                section="global",
                change_type="refresh",
                before="TOC/fields pending",
                after="TOC/fields refresh disabled by env",
                paragraph_index=-1,
                success=True,
                failure_reason=None,
            )

        # 5. 生成报告
        self._emit_progress(step_idx, total_steps, "正在生成处理报告…")
        self._generate_reports(output_paths)

        return self._finalize_result(
            doc=doc,
            original_doc=original_doc,
            output_paths=output_paths,
        )

    def _generate_reports(self, output_paths: dict) -> None:
        """生成 JSON / Markdown 报告文件（如果路径存在于 output_paths）。"""
        from src.report.collector import collect_report
        from src.report.json_report import generate_json_report
        from src.report.markdown_report import generate_markdown_report
        from src.engine.rules.base import ValidationIssue

        context_issues: list[ValidationIssue] = []
        for rec in self.tracker.records:
            if rec.rule_name == "validation":
                context_issues.append(ValidationIssue(
                    level="warning" if rec.success else "error",
                    rule_name=rec.rule_name,
                    message=rec.before,
                    location=rec.target,
                ))

        report_data = collect_report(
            self.tracker,
            scene_name=getattr(self.config, "name", ""),
            input_file="",
            validation_issues=context_issues,
        )

        json_path = output_paths.get("report_json")
        if json_path:
            try:
                generate_json_report(report_data, json_path)
            except Exception:
                pass

        md_path = output_paths.get("report_md")
        if md_path:
            try:
                generate_markdown_report(report_data, md_path)
            except Exception:
                pass

    @staticmethod
    def _find_keyword_index(doc: Document, keyword: str) -> int | None:
        """在文档中查找包含关键字的段落索引"""
        import re
        kw = re.sub(r'\s+', '', keyword)
        for i, para in enumerate(doc.paragraphs):
            text = re.sub(r'\s+', '', para.text.strip())
            if kw in text:
                return i
        return None

    @staticmethod
    def _find_page_start_index(doc: Document, target_page: int) -> int | None:
        """根据物理页码找到该页第一个段落的索引。

        通过统计显式分页符（hard page break）和分节符（section break）
        来估算每一页的起始段落。target_page 从 1 开始。

        注意：仅能追踪显式分页，无法感知 Word 因内容溢出而发生的
        自动分页（soft page break），因此对于长段落溢出的情况可能
        存在少量偏差。
        """
        from lxml import etree

        if target_page is None or target_page < 1:
            return None
        if target_page == 1:
            return 0

        _NSMAP = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        current_page = 1
        body = doc.element.body
        para_idx = -1

        for child in body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "p":
                para_idx += 1

                # 检查段落内是否有硬分页符：<w:br w:type="page"/>
                for br in child.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"):
                    br_type = br.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type", "")
                    if br_type == "page":
                        current_page += 1
                        if current_page == target_page:
                            # 分页符在当前段落内，下一段落才是新页的开始
                            # 但如果分页符是段落最后一个元素，该段实际属于前一页
                            return para_idx + 1

                # 检查段落的分节符（pPr/sectPr）
                pPr = child.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
                if pPr is not None:
                    sect = pPr.find(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
                    if sect is not None:
                        sect_type_el = sect.find(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type")
                        sect_type = ""
                        if sect_type_el is not None:
                            sect_type = sect_type_el.get(
                                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "")
                        # nextPage / oddPage / evenPage 会产生新页
                        if sect_type in ("nextPage", "oddPage", "evenPage", ""):
                            current_page += 1
                            if current_page == target_page:
                                return para_idx + 1

            elif tag == "sectPr":
                # 文档最后一个 sectPr（body 级别），通常不产生新页
                pass

        # 页码超出文档实际页数
        return None

    def _save_outputs(self, doc: Document, src_path: Path) -> dict:
        """保存输出文件，返回路径字典

        策略：原始文件保持不变，新生成文件以 _new 后缀保存在同目录。
        """
        stem = src_path.stem
        out_dir = src_path.parent
        sub_dir = out_dir / f"{stem}_排版附件"
        sub_dir.mkdir(exist_ok=True)
        paths = {"sub_dir": str(sub_dir)}

        # 原始文件保持不变，不做备份/重命名
        paths["original"] = str(src_path)

        # 最终稿：保存为 {stem}_new.docx（不覆盖原文件）
        if self.config.output.final_docx:
            final_path = self._save_doc_with_fallback(
                doc, out_dir / f"{stem}_new.docx"
            )
            paths["final"] = str(final_path)

        # 对比稿 → 子文件夹
        if self.config.output.compare_docx:
            paths["compare"] = str(sub_dir / f"{stem}_对比稿.docx")

        # 报告 → 子文件夹
        if self.config.output.report_json:
            paths["report_json"] = str(sub_dir / f"{stem}_报告.json")
        if self.config.output.report_markdown:
            paths["report_md"] = str(sub_dir / f"{stem}_报告.md")

        return paths

    @staticmethod
    def _save_doc_with_fallback(doc: Document, primary_path: Path) -> Path:
        """保存文档，若目标文件被占用则自动尝试递增后缀路径。"""
        try:
            doc.save(str(primary_path))
            return primary_path
        except PermissionError:
            pass

        base_dir = primary_path.parent
        stem = primary_path.stem
        suffix = primary_path.suffix
        for idx in range(2, 100):
            alt = base_dir / f"{stem}_{idx}{suffix}"
            try:
                doc.save(str(alt))
                return alt
            except PermissionError:
                continue

        # 连续失败时抛出原始异常语义
        doc.save(str(primary_path))
        return primary_path
